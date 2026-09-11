from __future__ import annotations
from pathlib import Path
import pymupdf as fitz
import docx  # python-docx


def parse_docx(path: Path) -> list[dict]:
    """DOCX has no fixed page concept - page breaks are a rendering-time
    property of fonts/margins, not stored positions. Treated as one logical
    unit (page=1), same convention already used for .txt/.md; chunk_pages
    still splits it by token count regardless."""
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    return [{"source": path.name, "page": 1, "text": text, "suspect_scanned": len(text) < 40}]


def parse(path: Path) -> list[dict]:
    """Page-level text extraction.

    pymupdf is roughly 50x faster than an `unstructured` hi_res pipeline and
    handles most born-digital PDFs. If a page yields almost no text the file is
    probably scanned and needs OCR - we flag it rather than silently returning
    empty chunks, because silent empty chunks are how RAG quality dies.
    """
    doc = fitz.open(path)
    pages = []
    for i, page in enumerate(doc):
        text = page.get_text("text").strip()
        pages.append({"source": path.name, "page": i + 1, "text": text,
                      "suspect_scanned": len(text) < 40})
    doc.close()
    return pages


def parse_dir(corpus: Path) -> list[dict]:
    out: list[dict] = []
    for p in sorted(corpus.rglob("*")):
        if p.suffix.lower() == ".pdf":
            out.extend(parse(p))
        elif p.suffix.lower() == ".docx":
            out.extend(parse_docx(p))
        elif p.suffix.lower() in {".txt", ".md"}:
            out.append({"source": p.name, "page": 1,
                        "text": p.read_text(encoding="utf-8", errors="ignore"),
                        "suspect_scanned": False})
    return out
